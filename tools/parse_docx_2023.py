# -*- coding: utf-8 -*-
"""Trích đề JLPT N3 tháng 12/2023 từ bản gõ lại .docx.

    cd tools && python parse_docx_2023.py

Viết riêng chứ không dùng chung với `parse_docx_2021.py` vì hai bản gõ lại có
cách trình bày khác hẳn nhau:

| | 2021-07 | 2023-12 |
|---|---|---|
| mốc phần thi | có 「第N部分」 | không có, nhận ra nhờ số 問題 quay về 1 |
| lựa chọn | mỗi dòng một cái | lúc bốn cái một dòng, lúc mỗi dòng một cái |
| số câu | `1.` và `1番` | `1.`, số trần `19`, và `1ばん` |
| bài đọc 問題7 | nằm trước câu hỏi | nằm **sau** câu hỏi |
| đáp án | có bảng ở cuối | **không có** |

Ảnh phần nghe trong file này là ảnh chụp cả trang, mỗi ảnh hai câu, nên phải
tách đôi theo khung tranh (`split_panels`).
"""
import io
import json
import os
import re
import sys

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import docx
from PIL import Image

here = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(os.path.expanduser("~"), "Downloads", "JLPT_N3_12_2023_FIXED.docx")
TAG = "2023-12"
TITLE = "Đề thi tháng 12/2023"
IMGDIR = os.path.join(here, "..", "images", TAG)
AUDIO = "../audio/%s/choukai.mp3" % TAG
OUT = os.path.join(here, "exam_2023_12.json")

SECTIONS = {
    "moji":    {"name": "Từ vựng – Chữ Hán", "jp": "言語知識（文字・語彙）", "minutes": 30},
    "bunpou":  {"name": "Ngữ pháp – Đọc hiểu", "jp": "言語知識（文法）・読解", "minutes": 70},
    "choukai": {"name": "Nghe hiểu", "jp": "聴解", "minutes": 40},
}
RX_MONDAI = re.compile(r"^問題[　 ]?(\d)[　 ]")
RX_Q = re.compile(r"^(\d{1,2})[\.．][　 ]*(.*)$")
RX_BARE = re.compile(r"^(\d{1,2})$")               # 問題3 văn bản: số câu đứng một mình
RX_BAN = re.compile(r"^(\d{1,2})ばん")
RX_OPT1 = re.compile(r"^([1-4])[　 ]+(\S.*)$")     # mỗi dòng một lựa chọn
SHARED_PASSAGE = {("bunpou", 3), ("bunpou", 5), ("bunpou", 6), ("bunpou", 7)}
PER_Q_PASSAGE = {("bunpou", 4)}
AFTER_Q_PASSAGE = {("bunpou", 7)}                  # bài đọc in sau câu hỏi
AUDIO_ONLY = {3: (3, 4), 5: (9, 3)}                # 問題: (số câu, số lựa chọn)
# đề in nhiều câu trên một trang ảnh -> mỗi ảnh mấy câu của 問題4
PANEL_LABELS = [(1, 2), (3, 4)]


def para_text(p):
    out = []
    for r in p.runs:
        out.append("【%s】" % r.text if r.underline and r.text.strip() else r.text)
    return re.sub(r"[ \t]+", " ", "".join(out)).strip()


def four_on_one_line(t):
    """'1 あ 2 い 3 う 4 え' -> bốn lựa chọn; không đúng dạng thì trả None."""
    pos = []
    frm = 0
    for n in range(1, 5):
        m = re.compile(r"(?:^|[\s　])%d[\s　]" % n).search(t, frm)
        if not m:
            return None
        pos.append(m)
        frm = m.end()
    out = []
    for k in range(4):
        s = pos[k].end()
        e = pos[k + 1].start() if k < 3 else len(t)
        out.append(t[s:e].strip())
    return out if all(out) else None


def split_panels(img_bytes, count):
    """Ảnh chụp cả trang -> cắt lấy `count` khung tranh theo khoảng trắng giữa chúng."""
    im = Image.open(io.BytesIO(img_bytes)).convert("L")
    w, h = im.size
    px = im.load()
    thr = max(4, w // 4 // 40)
    runs, start = [], None
    for y in range(h):
        dark = sum(1 for x in range(0, w, 4) if px[x, y] < 160)
        if dark > thr and start is None:
            start = y
        elif dark <= thr and start is not None:
            runs.append((start, y))
            start = None
    if start is not None:
        runs.append((start, h))
    merged = []                       # nối các khối sát nhau (khung tranh bị đứt dòng)
    for a, b in runs:
        if merged and a - merged[-1][1] < 30:
            merged[-1][1] = b
        else:
            merged.append([a, b])
    panels = sorted([r for r in merged if r[1] - r[0] > 300],
                    key=lambda r: r[1] - r[0], reverse=True)[:count]
    panels.sort()
    if len(panels) != count:
        raise SystemExit("ảnh: tìm được %d khung, cần %d" % (len(panels), count))
    pad = 12
    return [im.crop((0, max(0, a - pad), w, min(h, b + pad))) for a, b in panels]


def parse():
    doc = docx.Document(SRC)
    secs = []
    cur_sec = cur_m = cur_q = None
    buf = []
    imgs = {}
    for i, p in enumerate(doc.paragraphs):
        ids = re.findall(r'r:embed="([^"]+)"', p._p.xml)
        if ids:
            imgs[i] = doc.part.related_parts[ids[0]].blob

    def new_sec(key):
        s = dict(SECTIONS[key], key=key, mondai=[])
        secs.append(s)
        return s

    def flush(target_q):
        nonlocal buf
        if not buf or cur_m is None:
            buf = []
            return
        text = "\n".join(buf).strip()
        key = (cur_sec["key"], cur_m["no"])
        if key in PER_Q_PASSAGE and target_q is not None:
            target_q["passage"] = text
        else:
            cur_m["passage"] = (cur_m.get("passage", "") + "\n" + text).strip()
        buf = []

    pend_img = []           # ảnh của 問題4 phần nghe, xử lý sau khi biết 問題
    for i, p in enumerate(doc.paragraphs):
        t = para_text(p)
        if i in imgs:
            if cur_sec and cur_sec["key"] == "choukai" and cur_m and cur_m["no"] == 4:
                pend_img.append(imgs[i])
            elif cur_q is not None:
                name = "choukai-m%d-q%d.png" % (cur_m["no"], cur_q["label"])
                os.makedirs(IMGDIR, exist_ok=True)
                open(os.path.join(IMGDIR, name), "wb").write(imgs[i])
                cur_q["img"] = "../images/%s/%s" % (TAG, name)
        # dòng tiêu đề phần thi in giữa file, không phải nội dung đề
        if (not t or t.startswith("Bản chuyển sang text") or t.startswith("2023 年")
                or t in ("文法・読解", "言語知識（文字・語彙）", "言語知識（文法）・読解")):
            continue
        if t.startswith("聴解"):
            flush(cur_q)
            cur_sec = new_sec("choukai")
            cur_m = {"no": 1, "instruction": "問題1 では、まず質問を聞いてください。"
                     "それから話を聞いて、1から4の中から、最もよいものを一つえらんでください。",
                     "questions": []}
            cur_sec["mondai"].append(cur_m)
            cur_q = None
            continue
        if t.startswith("絵の問題"):
            flush(cur_q)
            cur_m = {"no": 4, "instruction": "問題4 では、えを見ながら質問を聞いてください。"
                     "やじるし（→）の人は何と言いますか。1から3の中から、"
                     "最もよいものを一つえらんでください。", "questions": []}
            cur_sec["mondai"].append(cur_m)
            cur_q = None
            continue
        m = RX_MONDAI.match(t)
        if m:
            flush(cur_q)
            no = int(m.group(1))
            if cur_sec is None:
                cur_sec = new_sec("moji")
            elif no == 1 and cur_sec["key"] == "moji":
                cur_sec = new_sec("bunpou")
            cur_m = {"no": no, "instruction": t, "questions": []}
            cur_sec["mondai"].append(cur_m)
            cur_q = None
            continue
        if cur_sec is None or cur_m is None:
            continue
        key = (cur_sec["key"], cur_m["no"])
        # lựa chọn
        if cur_q is not None and len(cur_q["opts"]) < 4:
            # thử bốn-cái-một-dòng TRƯỚC: dòng đó cũng khớp dạng một-lựa-chọn-một-dòng
            # nên xét sau là nuốt cả bốn vào làm một
            if not cur_q["opts"]:
                four = four_on_one_line(t)
                if four:
                    cur_q["opts"] = four
                    continue
            om = RX_OPT1.match(t)
            if om and int(om.group(1)) == len(cur_q["opts"]) + 1:
                cur_q["opts"].append(om.group(2).strip())
                continue
        # câu hỏi mới
        if cur_sec["key"] == "choukai":
            qm = RX_BAN.match(t)
            if qm:
                lb = int(qm.group(1))
                if lb == 1 and cur_m["questions"]:      # ばん quay về 1 = sang 問題 sau
                    cur_m = {"no": cur_m["no"] + 1,
                             "instruction": "問題2 では、まず質問を聞いてください。そのあと、"
                             "問題用紙を見てください。読む時間があります。それから話を聞いて、"
                             "1から4の中から、最もよいものを一つえらんでください。",
                             "questions": []}
                    cur_sec["mondai"].append(cur_m)
                cur_q = {"label": lb, "q": "", "opts": [], "answer": None, "audio": AUDIO}
                cur_m["questions"].append(cur_q)
                continue
        else:
            qm = RX_Q.match(t) or (RX_BARE.match(t) if key == ("bunpou", 3) else None)
            if qm:
                body = qm.group(2).strip() if qm.lastindex and qm.lastindex > 1 else ""
                cur_q = {"label": int(qm.group(1)), "q": body, "opts": [], "answer": None}
                flush(cur_q)
                cur_m["questions"].append(cur_q)
                continue
        # dòng tiếp của đề bài (hội thoại nhiều dòng)
        if cur_q is not None and not cur_q["opts"] and key not in AFTER_Q_PASSAGE:
            cur_q["q"] = (cur_q["q"] + " " + t).strip()
            continue
        buf.append(t)
    flush(cur_q)

    # bảng giá của 問題7 nằm trong table
    tbls = "\n".join("\n".join(" ｜ ".join(c.text.strip().replace("\n", " ") for c in r.cells)
                               for r in tb.rows) for tb in doc.tables)
    m7 = next(m for s in secs if s["key"] == "bunpou" for m in s["mondai"] if m["no"] == 7)
    m7["passage"] = (m7.get("passage", "") + "\n" + tbls).strip()

    cho = next(s for s in secs if s["key"] == "choukai")
    # 問題4: mỗi ảnh chụp là một trang gồm hai câu -> tách đôi
    m4 = next(m for m in cho["mondai"] if m["no"] == 4)
    os.makedirs(IMGDIR, exist_ok=True)
    for blob, labels in zip(pend_img, PANEL_LABELS):
        for panel, lb in zip(split_panels(blob, len(labels)), labels):
            name = "choukai-m4-q%d.png" % lb
            panel.save(os.path.join(IMGDIR, name), optimize=True)
            m4["questions"].append({
                "label": lb, "q": "", "opts": ["", "", ""], "answer": None,
                "audio": AUDIO, "img": "../images/%s/%s" % (TAG, name)})
    m4["questions"].sort(key=lambda q: q["label"])
    # 問題3 và 問題5: đề không in gì, vẫn tạo đủ số câu theo format JLPT
    for no, (count, nopt) in AUDIO_ONLY.items():
        mon = {"no": no, "instruction":
               "問題%d では、問題用紙に何もいんさつされていません。" % no, "questions": [
                   {"label": lb, "q": "", "opts": [""] * nopt, "answer": None,
                    "audio": AUDIO} for lb in range(1, count + 1)]}
        cho["mondai"].append(mon)
    cho["mondai"].sort(key=lambda m: m["no"])
    for m in cho["mondai"]:
        for q in m["questions"]:
            if not q["opts"]:
                q["opts"] = [""] * (3 if m["no"] in (4, 5) else 4)
    for s in secs:
        for m in s["mondai"]:
            if not m.get("passage"):
                m.pop("passage", None)
    return [{"id": TAG, "title": TITLE, "sections": secs}]


if __name__ == "__main__":
    exams = parse()
    tot = 0
    for s in exams[0]["sections"]:
        print("[%s] %s" % (s["key"], s["name"]))
        for m in s["mondai"]:
            n = len(m["questions"])
            tot += n
            bad = [q["label"] for q in m["questions"]
                   if len(q["opts"]) not in (3, 4) or any(
                       not str(o).strip() for o in q["opts"]) and any(
                       str(o).strip() for o in q["opts"])]
            print("   問題%d: %2d câu%s | %s" % (
                m["no"], n, "  LỖI lựa chọn: %s" % bad if bad else "",
                "passage %d ký tự" % len(m["passage"]) if m.get("passage") else "-"))
    print("\ntổng %d câu | ảnh: %d" % (tot, len(os.listdir(IMGDIR))))
    print("!! docx này KHÔNG có đáp án -> mọi câu answer=null, chưa chấm điểm được")
    tmp = OUT + ".tmp"
    with io.open(tmp, "w", encoding="utf-8") as f:
        json.dump(exams, f, ensure_ascii=False, indent=1)
        f.write("\n")
    os.replace(tmp, OUT)
    print("ghi: %s" % OUT)
